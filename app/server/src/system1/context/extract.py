"""Turns an uploaded file into plain text. The file is data: nothing in it is ever executed or followed."""

import io
import zipfile
from pathlib import PurePosixPath, PureWindowsPath

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import DependencyError, PyPdfError

from system1.errors import ApiError

MAX_NAME_CHARS = 200
_TEXT_KINDS = {".md": "md", ".markdown": "md", ".txt": "txt"}


class UnsupportedFile(ApiError):
    status_code = 422
    code = "unsupported_file"
    message = "Use a PDF, Word, Markdown or text file."


class EmptyText(ApiError):
    status_code = 422
    code = "empty_text"
    message = "We couldn't read any text from that file. Scanned PDFs aren't supported yet."


class UnreadablePdf(ApiError):
    status_code = 422
    code = "unreadable_pdf"
    message = "We couldn't read that PDF."


class UnreadableDocx(ApiError):
    status_code = 422
    code = "unreadable_docx"
    message = "We couldn't read that Word file."


def clean_name(filename: str) -> str:
    """Only the base name is kept: a client-supplied path never reaches storage or the prompt."""
    base = PureWindowsPath(PurePosixPath(filename.replace("\\", "/")).name).name.strip()
    return base[:MAX_NAME_CHARS]


def clean_text(text: str) -> str:
    return text.replace("\x00", "").strip()


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnreadablePdf()
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except (PyPdfError, DependencyError):
        raise UnreadablePdf() from None


def _docx_text(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError):
        raise UnreadableDocx() from None
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    name = clean_name(filename).lower()
    suffix = PurePosixPath(name).suffix
    if suffix == ".pdf":
        kind, text = "pdf", _pdf_text(data)
    elif suffix == ".docx":
        kind, text = "docx", _docx_text(data)
    elif suffix in _TEXT_KINDS:
        kind, text = _TEXT_KINDS[suffix], data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedFile()
    text = clean_text(text)
    if not text:
        raise EmptyText()
    return kind, text
